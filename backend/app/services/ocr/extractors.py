import os
import shutil
import subprocess
import tempfile

from PIL import Image
from pypdf import PdfReader

try:
    from docx import Document
except Exception:
    Document = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    import xlrd
except Exception:
    xlrd = None

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from pdf2image import convert_from_path
except Exception:
    convert_from_path = None

from app.utils.file_utils import decode_bytes_safely, csv_text_from_bytes
from .image_utils import prepare_base_image, preprocess_image_variants_for_ocr
from .ocr_engine import choose_best_ocr_text, run_ocr_with_multiple_configs
from .lab_parser import build_cbc_panel_json, build_compact_cbc_text
from .text_utils import (
    extract_structured_lab_lines,
    normalize_ocr_text,
    postprocess_extracted_text,
    score_ocr_text,
)


def build_final_llm_text(raw_text: str) -> str:
    structured = extract_structured_lab_lines(raw_text)
    compact = build_compact_cbc_text(structured)

    if compact.strip():
        return compact.strip()

    return raw_text.strip()


def extract_text_from_image(file_path: str) -> str:
    image = prepare_base_image(file_path)

    candidates = []
    candidates.extend(run_ocr_with_multiple_configs(image))

    for processed_image in preprocess_image_variants_for_ocr(image):
        candidates.extend(run_ocr_with_multiple_configs(processed_image))

    print(f"OCR candidates total: {len(candidates)}")

    best_text = choose_best_ocr_text(candidates)

    print("BEST OCR TEXT START")
    print(best_text[:4000])
    print("BEST OCR TEXT END")

    final_text = postprocess_extracted_text(best_text)

    print("POSTPROCESSED OCR TEXT START")
    print(final_text[:4000])
    print("POSTPROCESSED OCR TEXT END")

    structured = extract_structured_lab_lines(final_text)

    print("STRUCTURED LAB LINES START")
    for item in structured:
        print(item)
    print("STRUCTURED LAB LINES END")

    parsed_json = build_cbc_panel_json(structured)
    print("PARSED CBC JSON START")
    print(parsed_json)
    print("PARSED CBC JSON END")

    compact_text = build_final_llm_text(final_text)
    print("COMPACT LLM TEXT START")
    print(compact_text[:4000])
    print("COMPACT LLM TEXT END")

    return compact_text.strip()


def extract_text_from_pdf_direct(file_path: str) -> str:
    text_parts = []

    try:
        reader = PdfReader(file_path)

        for page in reader.pages:
            page_text = page.extract_text() or ""
            page_text = normalize_ocr_text(page_text)

            if page_text.strip():
                text_parts.append(page_text.strip())
    except Exception:
        pass

    cleaned = postprocess_extracted_text("\n".join(text_parts).strip())
    return build_final_llm_text(cleaned)


def extract_text_from_pdf_via_ocr(file_path: str) -> str:
    if convert_from_path is None:
        return ""

    try:
        images = convert_from_path(file_path)
    except Exception:
        return ""

    ocr_parts = []

    for page_index, img in enumerate(images, start=1):
        if not isinstance(img, Image.Image):
            continue

        page_candidates = []
        page_candidates.extend(run_ocr_with_multiple_configs(img))

        for processed_img in preprocess_image_variants_for_ocr(img):
            page_candidates.extend(run_ocr_with_multiple_configs(processed_img))

        print(f"PDF page {page_index}: OCR candidates total = {len(page_candidates)}")

        best_page_text = choose_best_ocr_text(page_candidates)

        print(f"PDF page {page_index}: BEST OCR TEXT START")
        print(best_page_text[:4000])
        print(f"PDF page {page_index}: BEST OCR TEXT END")

        final_page_text = postprocess_extracted_text(best_page_text)

        print(f"PDF page {page_index}: POSTPROCESSED OCR TEXT START")
        print(final_page_text[:4000])
        print(f"PDF page {page_index}: POSTPROCESSED OCR TEXT END")

        if final_page_text.strip():
            ocr_parts.append(final_page_text.strip())

    cleaned = postprocess_extracted_text("\n".join(ocr_parts).strip())
    return build_final_llm_text(cleaned)


def extract_text_from_pdf(file_path: str) -> str:
    direct_text = extract_text_from_pdf_direct(file_path)
    ocr_text = extract_text_from_pdf_via_ocr(file_path)

    if score_ocr_text(ocr_text) > score_ocr_text(direct_text):
        return ocr_text

    return direct_text


def extract_text_from_docx(file_path: str) -> str:
    if Document is None:
        return ""

    try:
        doc = Document(file_path)
        parts = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_values = []

                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_values.append(cell_text)

                if row_values:
                    parts.append(" | ".join(row_values))

        text = "\n".join(parts).strip()
        cleaned = postprocess_extracted_text(text)
        return build_final_llm_text(cleaned)

    except Exception:
        return ""


def extract_text_from_doc(file_path: str) -> str:
    libreoffice_path = shutil.which("soffice") or shutil.which("libreoffice")

    if not libreoffice_path:
        return ""

    temp_dir = tempfile.mkdtemp(prefix="doc_convert_")

    try:
        subprocess.run(
            [
                libreoffice_path,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                temp_dir,
                file_path,
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

        base_name = os.path.splitext(os.path.basename(file_path))[0]
        converted_docx_path = os.path.join(temp_dir, f"{base_name}.docx")

        if os.path.exists(converted_docx_path):
            return extract_text_from_docx(converted_docx_path)

        return ""

    except Exception:
        return ""

    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def extract_text_from_excel(file_path: str, file_extension: str) -> str:
    parts = []

    if pd is not None:
        try:
            excel_file = pd.ExcelFile(file_path)

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)

                parts.append(f"[Лист: {sheet_name}]")
                parts.append(df.fillna("").astype(str).to_string(index=False, header=False))

            text = "\n".join(parts).strip()
            cleaned = postprocess_extracted_text(text)
            return build_final_llm_text(cleaned)

        except Exception:
            pass

    if file_extension == ".xlsx" and openpyxl is not None:
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)

            for sheet in workbook.worksheets:
                parts.append(f"[Лист: {sheet.title}]")

                for row in sheet.iter_rows(values_only=True):
                    row_values = [str(cell).strip() for cell in row if cell is not None]
                    if row_values:
                        parts.append(" | ".join(row_values))

            text = "\n".join(parts).strip()
            cleaned = postprocess_extracted_text(text)
            return build_final_llm_text(cleaned)

        except Exception:
            pass

    if file_extension == ".xls" and xlrd is not None:
        try:
            workbook = xlrd.open_workbook(file_path)

            for i in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(i)
                parts.append(f"[Лист: {sheet.name}]")

                for row_idx in range(sheet.nrows):
                    row_values = [
                        str(sheet.cell_value(row_idx, col_idx)).strip()
                        for col_idx in range(sheet.ncols)
                    ]
                    row_values = [v for v in row_values if v]

                    if row_values:
                        parts.append(" | ".join(row_values))

            text = "\n".join(parts).strip()
            cleaned = postprocess_extracted_text(text)
            return build_final_llm_text(cleaned)

        except Exception:
            pass

    return ""


def extract_text_from_csv(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw_data = f.read()

    text = csv_text_from_bytes(raw_data)
    cleaned = postprocess_extracted_text(text)
    return build_final_llm_text(cleaned)


def extract_text_from_txt(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw_data = f.read()

    text = decode_bytes_safely(raw_data).strip()
    cleaned = postprocess_extracted_text(text)
    return build_final_llm_text(cleaned)


def extract_text_from_file(file_path: str, file_extension: str) -> str:
    file_extension = file_extension.lower()

    if file_extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if file_extension in {".jpg", ".jpeg", ".png"}:
        return extract_text_from_image(file_path)

    if file_extension == ".docx":
        return extract_text_from_docx(file_path)

    if file_extension == ".doc":
        return extract_text_from_doc(file_path)

    if file_extension in {".xls", ".xlsx"}:
        return extract_text_from_excel(file_path, file_extension)

    if file_extension == ".csv":
        return extract_text_from_csv(file_path)

    if file_extension == ".txt":
        return extract_text_from_txt(file_path)

    return ""