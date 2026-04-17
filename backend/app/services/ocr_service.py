
# ИМПОРТЫ СТАНДАРТНОЙ БИБЛИОТЕКИ PYTHON
# os нужен для работы с путями к файлам:
# - получить имя файла
# - убрать расширение
# - проверить существование файла
import os
# shutil нужен, чтобы:
# - найти установленную программу в системе (например LibreOffice)
# - удалить временную папку целиком
import shutil
# subprocess нужен для запуска внешней программы из Python.
# В нашем случае он используется для запуска LibreOffice,
# чтобы конвертировать старый .doc в .docx
import subprocess
# tempfile нужен для создания временных папок,
# например во время конвертации .doc -> .docx
import tempfile
# ИМПОРТЫ СТОРОННИХ PYTHON-БИБЛИОТЕК
# pytesseract — это библиотека, которая позволяет
# использовать Tesseract OCR из Python.
# С её помощью мы распознаём текст на картинках.
import pytesseract

# PIL (Pillow) — библиотека для работы с изображениями.
# Мы используем:
# - Image для открытия изображения
# - ImageEnhance для повышения контраста
# - ImageFilter для повышения резкости
# - ImageOps для автоконтраста
from PIL import Image, ImageEnhance, ImageFilter, ImageOps

# PdfReader из pypdf нужен для чтения PDF-файлов
# и попытки вытащить текст напрямую из PDF,
# если PDF содержит текстовый слой.
from pypdf import PdfReader


# ============================================================
# ИМПОРТ НАШИХ ВСПОМОГАТЕЛЬНЫХ ФУНКЦИЙ
# ============================================================

# Эти функции лежат в нашем проекте в app/utils/file_utils.py
# decode_bytes_safely — безопасно декодирует байты в текст
# csv_text_from_bytes — превращает CSV в удобный текст
from app.utils.file_utils import decode_bytes_safely, csv_text_from_bytes


# ============================================================
# ОПЦИОНАЛЬНЫЕ ИМПОРТЫ
# ============================================================
# Ниже идут библиотеки, которые МОГУТ быть не установлены.
#
# Почему это сделано через try/except:
# чтобы backend не падал целиком, если какого-то модуля нет.
#
# Вместо ошибки приложение просто будет использовать
# более ограниченный режим работы.
# ============================================================

try:
    # python-docx нужен для чтения .docx файлов
    from docx import Document
except Exception:
    # Если библиотека не установлена, ставим None.
    # Это значит: DOCX сейчас не поддерживается полноценно.
    Document = None

try:
    # openpyxl используется для чтения .xlsx
    import openpyxl
except Exception:
    openpyxl = None

try:
    # xlrd используется для чтения старых .xls
    import xlrd
except Exception:
    xlrd = None

try:
    # pandas — удобный способ читать Excel-таблицы
    import pandas as pd
except Exception:
    pd = None

try:
    # pdf2image позволяет превратить PDF-страницы в картинки,
    # чтобы потом распознать их через OCR.
    from pdf2image import convert_from_path
except Exception:
    convert_from_path = None


# ============================================================
# ПРЕДОБРАБОТКА ИЗОБРАЖЕНИЯ ДЛЯ OCR
# ============================================================

def preprocess_image_for_ocr(image: Image.Image) -> Image.Image:
    """
    Эта функция улучшает изображение перед OCR-распознаванием.

    Зачем это нужно:
    Tesseract лучше распознаёт текст,
    если изображение предварительно "почистить".

    Что мы делаем:
    1. Переводим картинку в оттенки серого
    2. Улучшаем контраст
    3. Повышаем резкость
    4. Делаем чёрно-белую бинаризацию

    Это особенно полезно для:
    - плохих фото анализов
    - сканов
    - документов с тусклым текстом
    """

    # Переводим изображение в чёрно-белый серый режим.
    image = image.convert("L")

    # Автоматически улучшаем контраст.
    image = ImageOps.autocontrast(image)

    # Добавляем резкость,
    # чтобы контуры букв стали чуть лучше видны.
    image = image.filter(ImageFilter.SHARPEN)

    # Повышаем контраст ещё сильнее.
    enhancer = ImageEnhance.Contrast(image)
    image = enhancer.enhance(1.5)

    # Бинаризация:
    # всё темнее порога делаем чёрным,
    # всё светлее — белым.
    #
    # Это помогает OCR лучше отделять текст от фона.
    image = image.point(lambda x: 0 if x < 160 else 255, mode="1")

    return image


# ============================================================
# OCR ИЗ ИЗОБРАЖЕНИЯ
# ============================================================

def extract_text_from_image(file_path: str) -> str:
    """
    Извлекает текст из изображения через OCR.

    Как работает:
    1. Открываем изображение
    2. Делаем улучшенную версию для OCR
    3. Пытаемся распознать текст:
       - из оригинального изображения
       - из обработанного изображения
    4. Берём тот вариант, где текста получилось больше

    Это полезно, потому что:
    иногда исходная картинка распознаётся лучше,
    а иногда — обработанная.
    """

    # Открываем изображение по пути.
    image = Image.open(file_path)

    # Создаём улучшенную версию изображения.
    processed_image = preprocess_image_for_ocr(image)

    # OCR исходной картинки.
    raw_text = pytesseract.image_to_string(
        image,
        lang="rus+eng",
        config="--psm 6",
    )

    # OCR обработанной картинки.
    processed_text = pytesseract.image_to_string(
        processed_image,
        lang="rus+eng",
        config="--psm 6",
    )

    # Выбираем тот вариант, где текста больше.
    best_text = processed_text if len(processed_text) >= len(raw_text) else raw_text

    # Убираем лишние пробелы по краям.
    return best_text.strip()


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ PDF
# ============================================================

def extract_text_from_pdf(file_path: str) -> str:
    """
    Извлекает текст из PDF-файла.

    Логика такая:
    1. Сначала пытаемся достать текст напрямую через PdfReader
       (это работает, если PDF содержит нормальный текстовый слой).
    2. Если текста мало или нет —
       пытаемся превратить страницы PDF в изображения
       и распознать их через OCR.

    Это важно, потому что PDF бывают двух типов:
    - текстовые (текст можно вытащить сразу)
    - сканы/фото внутри PDF (тогда нужен OCR)
    """

    # Сюда будем складывать текст по страницам.
    text_parts = []

    # --------------------------------------------
    # Шаг 1. Пробуем обычное извлечение текста.
    # --------------------------------------------
    try:
        reader = PdfReader(file_path)

        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text.strip())
    except Exception:
        # Если не получилось — просто идём дальше.
        # Ошибка здесь не означает, что PDF безнадёжен,
        # возможно, его получится распознать через OCR.
        pass

    extracted_text = "\n".join(text_parts).strip()

    # Если текста уже достаточно — возвращаем его сразу.
    if len(extracted_text) > 80:
        return extracted_text

    # --------------------------------------------
    # Шаг 2. Если текста мало, пробуем OCR PDF.
    # --------------------------------------------
    if convert_from_path is not None:
        try:
            # Конвертируем страницы PDF в изображения.
            images = convert_from_path(file_path)

            ocr_parts = []

            # Проходим по каждой странице-картинке.
            for img in images:
                processed_img = preprocess_image_for_ocr(img)

                page_text = pytesseract.image_to_string(
                    processed_img,
                    lang="rus+eng",
                    config="--psm 6",
                )

                if page_text.strip():
                    ocr_parts.append(page_text.strip())

            ocr_text = "\n".join(ocr_parts).strip()

            # Если OCR дал текста больше, чем обычный способ,
            # возвращаем именно OCR-вариант.
            if len(ocr_text) > len(extracted_text):
                return ocr_text

        except Exception:
            # Если OCR по PDF не сработал — просто вернём то, что уже есть.
            pass

    return extracted_text


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ DOCX
# ============================================================

def extract_text_from_docx(file_path: str) -> str:
    """
    Извлекает текст из файла DOCX.

    Что читаем:
    - обычные абзацы
    - таблицы

    Это важно, потому что медицинские документы часто содержат
    данные именно в таблицах.
    """

    # Если python-docx не установлен, DOCX сейчас обработать не можем.
    if Document is None:
        return ""

    try:
        # Открываем DOCX-файл.
        doc = Document(file_path)

        # Сюда будем складывать части текста.
        parts = []

        # --------------------------------------------
        # Читаем обычные абзацы.
        # --------------------------------------------
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        # --------------------------------------------
        # Читаем таблицы.
        # --------------------------------------------
        for table in doc.tables:
            for row in table.rows:
                row_values = []

                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_values.append(cell_text)

                if row_values:
                    # Собираем строку таблицы в формат:
                    # ячейка1 | ячейка2 | ячейка3
                    parts.append(" | ".join(row_values))

        return "\n".join(parts).strip()

    except Exception:
        return ""


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ СТАРОГО DOC
# ============================================================

def extract_text_from_doc(file_path: str) -> str:
    """
    Извлекает текст из старого формата .doc.

    Проблема:
    .doc — старый бинарный формат, его неудобно читать напрямую.

    Поэтому логика такая:
    1. Ищем LibreOffice в системе
    2. Если он есть — конвертируем .doc в .docx
    3. Потом читаем уже .docx

    Если LibreOffice не установлен,
    вернуть текст не сможем.
    """

    # Ищем путь к LibreOffice / soffice.
    libreoffice_path = shutil.which("soffice") or shutil.which("libreoffice")

    # Если LibreOffice в системе нет — обработка .doc невозможна.
    if not libreoffice_path:
        return ""

    # Создаём временную папку для результата конвертации.
    temp_dir = tempfile.mkdtemp(prefix="doc_convert_")

    try:
        # Запускаем LibreOffice в headless-режиме (без интерфейса),
        # чтобы сконвертировать файл в docx.
        subprocess.run(
            [
                libreoffice_path,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                temp_dir,
                file_path,
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

        # Получаем базовое имя файла без расширения.
        base_name = os.path.splitext(os.path.basename(file_path))[0]

        # Ожидаемый путь до нового docx.
        converted_docx_path = os.path.join(temp_dir, f"{base_name}.docx")

        # Если файл реально появился — читаем его как DOCX.
        if os.path.exists(converted_docx_path):
            return extract_text_from_docx(converted_docx_path)

        return ""

    except Exception:
        return ""

    finally:
        # В любом случае удаляем временную папку,
        # чтобы не засорять систему.
        shutil.rmtree(temp_dir, ignore_errors=True)


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ EXCEL
# ============================================================

def extract_text_from_excel(file_path: str, file_extension: str) -> str:
    """
    Извлекает текст из Excel-файлов.

    Поддерживаем:
    - .xlsx
    - .xls

    Логика:
    1. Сначала пробуем через pandas (удобнее всего)
    2. Если не получилось:
       - для .xlsx пробуем openpyxl
       - для .xls пробуем xlrd
    """

    # Сюда складываем итоговый текст.
    parts = []

    # --------------------------------------------
    # Шаг 1. Пробуем через pandas
    # --------------------------------------------
    if pd is not None:
        try:
            excel_file = pd.ExcelFile(file_path)

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)

                # Добавляем заголовок листа
                parts.append(f"[Лист: {sheet_name}]")

                # Превращаем таблицу в текст
                parts.append(
                    df.fillna("").astype(str).to_string(index=False, header=False)
                )

            text = "\n".join(parts).strip()
            if text:
                return text

        except Exception:
            pass

    # --------------------------------------------
    # Шаг 2. Если это xlsx, пробуем openpyxl
    # --------------------------------------------
    if file_extension == ".xlsx" and openpyxl is not None:
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)

            for sheet in workbook.worksheets:
                parts.append(f"[Лист: {sheet.title}]")

                for row in sheet.iter_rows(values_only=True):
                    row_values = [str(cell).strip() for cell in row if cell is not None]

                    if row_values:
                        parts.append(" | ".join(row_values))

            return "\n".join(parts).strip()

        except Exception:
            pass

    # --------------------------------------------
    # Шаг 3. Если это xls, пробуем xlrd
    # --------------------------------------------
    if file_extension == ".xls" and xlrd is not None:
        try:
            workbook = xlrd.open_workbook(file_path)

            for i in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(i)
                parts.append(f"[Лист: {sheet.name}]")

                for row_idx in range(sheet.nrows):
                    row_values = [
                        str(sheet.cell_value(row_idx, col_idx)).strip()
                        for col_idx in range(sheet.ncols)
                    ]

                    row_values = [v for v in row_values if v]

                    if row_values:
                        parts.append(" | ".join(row_values))

            return "\n".join(parts).strip()

        except Exception:
            pass

    return ""


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ CSV
# ============================================================

def extract_text_from_csv(file_path: str) -> str:
    """
    Читает CSV-файл и превращает его в обычный текст.

    Это удобно, потому что дальше AI проще анализировать текст,
    чем сырую табличную структуру.
    """

    # Открываем CSV как байты.
    with open(file_path, "rb") as f:
        raw_data = f.read()

    # Передаём байты во вспомогательную функцию,
    # которая аккуратно декодирует и форматирует CSV.
    return csv_text_from_bytes(raw_data)


# ============================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ TXT
# ============================================================

def extract_text_from_txt(file_path: str) -> str:
    """
    Читает обычный txt-файл.

    decode_bytes_safely нужен, чтобы не упасть на кодировках.
    """

    with open(file_path, "rb") as f:
        raw_data = f.read()

    return decode_bytes_safely(raw_data).strip()


# ============================================================
# ГЛАВНАЯ ФУНКЦИЯ МАРШРУТИЗАЦИИ ПО ТИПУ ФАЙЛА
# ============================================================

def extract_text_from_file(file_path: str, file_extension: str) -> str:
    """
    Главная функция, которая выбирает нужный способ
    извлечения текста в зависимости от расширения файла.

    Пример:
    - .pdf  -> extract_text_from_pdf
    - .jpg  -> extract_text_from_image
    - .docx -> extract_text_from_docx
    - .xlsx -> extract_text_from_excel
    """

    if file_extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if file_extension in {".jpg", ".jpeg", ".png"}:
        return extract_text_from_image(file_path)

    if file_extension == ".docx":
        return extract_text_from_docx(file_path)

    if file_extension == ".doc":
        return extract_text_from_doc(file_path)

    if file_extension in {".xls", ".xlsx"}:
        return extract_text_from_excel(file_path, file_extension)

    if file_extension == ".csv":
        return extract_text_from_csv(file_path)

    if file_extension == ".txt":
        return extract_text_from_txt(file_path)

    # Если расширение не поддержано — возвращаем пустую строку.
    return ""